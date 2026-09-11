from __future__ import annotations

import hashlib
import importlib.util
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from scripts.cleaning_engine import CleaningOptions, UnsafeStructureError, build_plan, execute_plan
from scripts.file_adapters import InvalidInputError, SheetSelectionRequired, load_input, write_output
from scripts.reporting import UPGRADE_MESSAGE, render_report

try:
    from scripts.release_check import check_zip_contents
except ModuleNotFoundError:  # release_check.py is a development validator, not a Skill runtime file
    release_check_path = Path(__file__).resolve().parents[3] / "scripts" / "release_check.py"
    release_check_spec = importlib.util.spec_from_file_location("zeyang_release_check", release_check_path)
    assert release_check_spec is not None and release_check_spec.loader is not None
    release_check_module = importlib.util.module_from_spec(release_check_spec)
    release_check_spec.loader.exec_module(release_check_module)
    check_zip_contents = release_check_module.check_zip_contents


UPGRADE_PROMO = (
    "更多完整功能：如需多表合并、两表比对、表格拆分、复杂去重、汇总统计或组合任务，"
    "可查看泽洋更多 Excel 工具：[https://skillpay.alipay.com/public/zeyang](https://skillpay.alipay.com/public/zeyang)"
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def test_csv_cleaning_removes_empty_rows_columns_and_duplicates(tmp_path: Path) -> None:
    source = tmp_path / "dirty.csv"
    source.write_text(
        "姓名,金额,空列\n 张三 ,100,\n,,\n张三,100,\n李四,200,\n",
        encoding="utf-8-sig",
    )
    table = load_input(source)
    plan = build_plan(table, CleaningOptions(dedupe=True))
    result = execute_plan(plan)

    assert result.headers == ["姓名", "金额"]
    assert result.rows == [["张三", "100"], ["李四", "200"]]
    assert result.stats.removed_empty_rows == 1
    assert result.stats.removed_empty_columns == 1
    assert result.stats.removed_duplicate_rows == 1
    assert result.structural_changed is True


def test_csv_gbk_is_read_and_output_is_utf8_bom(tmp_path: Path) -> None:
    source = tmp_path / "gbk.csv"
    source.write_bytes("姓名,备注\n张三, 需要回访 \n".encode("gb18030"))
    output = tmp_path / "cleaned.csv"
    table = load_input(source)
    result = execute_plan(build_plan(table, CleaningOptions()))
    write_output(table, output, result.headers, result.rows, result.structural_changed)

    assert output.read_bytes().startswith(b"\xef\xbb\xbf")
    assert load_input(output).rows == [["张三", "需要回访"]]


def test_csv_quoted_multiline_field_preserves_newline(tmp_path: Path) -> None:
    source = tmp_path / "multiline.csv"
    source.write_text('name,note\nA,"line1\nline2"\nB,ok\n', encoding="utf-8")

    table = load_input(source)

    assert table.rows[0] == ["A", "line1\nline2"]


@pytest.mark.parametrize(
    ("token", "content", "expected"),
    [
        ("comma", "a,b\n1,2\n", ","),
        ("tab", "a\tb\n1\t2\n", "\t"),
        ("semicolon", "a;b\n1;2\n", ";"),
    ],
)
def test_csv_standard_delimiters_are_detected(tmp_path: Path, token: str, content: str, expected: str) -> None:
    source = tmp_path / f"{token}.csv"
    source.write_text(content, encoding="utf-8")

    table = load_input(source)

    assert table.delimiter == expected
    assert table.headers == ["a", "b"]
    assert table.rows == [["1", "2"]]


def test_csv_single_column_comma_is_blocked_as_ambiguous(tmp_path: Path) -> None:
    source = tmp_path / "ambiguous.csv"
    source.write_text("note\nBeijing,Chaoyang\nShanghai,Pudong\n", encoding="utf-8")

    with pytest.raises(InvalidInputError, match="--delimiter"):
        load_input(source)


def test_csv_explicit_delimiter_resolves_single_column_comma(tmp_path: Path) -> None:
    source = tmp_path / "single-column.csv"
    source.write_text("note\nBeijing,Chaoyang\nShanghai,Pudong\n", encoding="utf-8")

    table = load_input(source, delimiter="tab")

    assert table.delimiter == "\t"
    assert table.headers == ["note"]
    assert table.rows == [["Beijing,Chaoyang"], ["Shanghai,Pudong"]]


def test_partially_empty_column_is_preserved(tmp_path: Path) -> None:
    source = tmp_path / "partial.csv"
    source.write_text("A,B,C\n1,,x\n2,has-value,y\n", encoding="utf-8")
    result = execute_plan(build_plan(load_input(source), CleaningOptions()))

    assert result.headers == ["A", "B", "C"]
    assert result.stats.removed_empty_columns == 0


def test_header_only_file_preserves_headers(tmp_path: Path) -> None:
    source = tmp_path / "header-only.csv"
    source.write_text("姓名,部门\n", encoding="utf-8")
    result = execute_plan(build_plan(load_input(source), CleaningOptions()))

    assert result.headers == ["姓名", "部门"]
    assert result.rows == []
    assert result.stats.removed_empty_columns == 0
    assert any("仅有表头" in warning for warning in result.warnings)


def test_empty_file_and_unsupported_format_are_rejected(tmp_path: Path) -> None:
    empty = tmp_path / "empty.csv"
    empty.write_bytes(b"")
    with pytest.raises(InvalidInputError):
        load_input(empty)

    unsupported = tmp_path / "legacy.xls"
    unsupported.write_bytes(b"not an xls")
    with pytest.raises(Exception, match="不支持"):
        load_input(unsupported)


def test_internal_blank_csv_column_gets_template_protection(tmp_path: Path) -> None:
    source = tmp_path / "layout.csv"
    source.write_text("A,,C\n1,,3\n2,,4\n", encoding="utf-8")
    result = execute_plan(build_plan(load_input(source), CleaningOptions()))

    assert result.headers == ["A", "", "C"]
    assert result.stats.removed_empty_columns == 0
    assert result.template_protected is True
    assert any("版式" in warning for warning in result.warnings)


def _make_workbook(path: Path) -> None:
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["姓名", "部门", "空列"])
    data.append([" 张三 ", "销售", None])
    data.append([None, None, None])
    data.append(["李四", "运营", None])
    notes = workbook.create_sheet("Notes")
    notes.append(["保留", "原样"])
    workbook.save(path)


def test_xlsx_requires_sheet_when_multiple_nonempty_sheets_exist(tmp_path: Path) -> None:
    source = tmp_path / "multi.xlsx"
    _make_workbook(source)
    with pytest.raises(SheetSelectionRequired):
        load_input(source)


def test_xlsx_target_sheet_changes_and_other_sheet_is_preserved(tmp_path: Path) -> None:
    source = tmp_path / "multi.xlsx"
    output = tmp_path / "cleaned.xlsx"
    _make_workbook(source)
    table = load_input(source, "Data")
    result = execute_plan(build_plan(table, CleaningOptions()))
    write_output(table, output, result.headers, result.rows, result.structural_changed)

    workbook = load_workbook(output, data_only=False)
    assert workbook["Notes"]["A1"].value == "保留"
    assert workbook["Data"]["A2"].value == "张三"
    assert workbook["Data"].max_column == 2
    assert workbook["Data"].max_row == 3


def test_xlsx_explicit_column_width_does_not_trigger_template_protection(tmp_path: Path) -> None:
    source = tmp_path / "width.xlsx"
    output = tmp_path / "width_cleaned.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Data"
    ws.column_dimensions["A"].width = 12
    ws.row_dimensions[2].height = 18
    ws.append(["姓名", "空列"])
    ws.append(["张三", None])
    ws.append([None, None])
    ws.append(["李四", None])
    workbook.save(source)

    table = load_input(source, "Data")
    result = execute_plan(build_plan(table, CleaningOptions()))
    write_output(
        table,
        output,
        result.headers,
        result.rows,
        result.structural_changed,
        result.source_row_indices,
        result.source_column_indices,
    )

    assert result.template_protected is False
    assert result.stats.removed_empty_rows == 1
    assert result.stats.removed_empty_columns == 1
    cleaned = load_workbook(output, data_only=False)["Data"]
    assert cleaned.max_row == 3
    assert cleaned.max_column == 1


def test_version_matches_pyproject() -> None:
    version = Path("VERSION").read_text(encoding="utf-8").strip()
    pyproject = Path("pyproject.toml").read_text(encoding="utf-8")
    match = re.search(r"(?m)^\s*version\s*=\s*[\"']([^\"']+)[\"']\s*$", pyproject)

    assert match is not None
    assert version.removeprefix("V") == match.group(1)


def test_xlsx_number_format_is_preserved_after_structural_cleaning(tmp_path: Path) -> None:
    source = tmp_path / "number_format.xlsx"
    output = tmp_path / "number_format_cleaned.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Data"
    ws.append(["金额", "空列"])
    ws.append([1200.5, None])
    ws["A2"].number_format = "0.00"
    ws.append([None, None])
    workbook.save(source)

    table = load_input(source, "Data")
    result = execute_plan(build_plan(table, CleaningOptions()))
    write_output(
        table,
        output,
        result.headers,
        result.rows,
        result.structural_changed,
        result.source_row_indices,
        result.source_column_indices,
    )

    cleaned = load_workbook(output, data_only=False)["Data"]
    assert cleaned["A2"].value == 1200.5
    assert cleaned["A2"].number_format == "0.00"


def test_merged_cells_are_protected(tmp_path: Path) -> None:
    source = tmp_path / "merged.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Data"
    ws.merge_cells("A1:B1")
    ws["A1"] = "模板标题"
    ws.append(["姓名", "部门"])
    workbook.save(source)

    table = load_input(source, "Data")
    with pytest.raises(UnsafeStructureError, match="合并单元格"):
        build_plan(table, CleaningOptions())


def test_formula_structure_is_not_deleted(tmp_path: Path) -> None:
    source = tmp_path / "formula.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Data"
    ws.append(["金额", "含税"])
    ws.append([100, "=A2*1.13"])
    workbook.save(source)

    table = load_input(source, "Data")
    result = execute_plan(build_plan(table, CleaningOptions()))

    assert result.structural_changed is False
    assert result.rows == [[100, "=A2*1.13"]]


def test_formula_table_with_actual_empty_row_change_is_blocked(tmp_path: Path) -> None:
    source = tmp_path / "formula_empty_row.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Data"
    ws.append(["金额", "含税"])
    ws.append([100, "=A2*1.13"])
    ws.append([None, None])
    ws.append([200, "=A4*1.13"])
    workbook.save(source)

    table = load_input(source, "Data")
    with pytest.raises(UnsafeStructureError, match="实际需要改变"):
        build_plan(table, CleaningOptions())


def test_formula_table_with_actual_empty_column_change_is_blocked(tmp_path: Path) -> None:
    source = tmp_path / "formula_empty_column.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Data"
    ws.append(["金额", "含税", "空列"])
    ws.append([100, "=A2*1.13", None])
    ws.append([200, "=A3*1.13", None])
    workbook.save(source)

    table = load_input(source, "Data")
    with pytest.raises(UnsafeStructureError, match="实际需要改变"):
        build_plan(table, CleaningOptions())


def test_formula_table_with_actual_duplicate_removal_is_blocked(tmp_path: Path) -> None:
    source = tmp_path / "formula_duplicate.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Data"
    ws.append(["金额", "含税"])
    ws.append([100, "=A2*1.13"])
    ws.append([100, "=A2*1.13"])
    workbook.save(source)

    table = load_input(source, "Data")
    with pytest.raises(UnsafeStructureError, match="实际需要改变"):
        build_plan(table, CleaningOptions(dedupe=True))


def test_duplicate_headers_can_be_renamed_and_reported(tmp_path: Path) -> None:
    source = tmp_path / "headers.csv"
    source.write_text("金额,金额\n1,2\n", encoding="utf-8")
    result = execute_plan(
        build_plan(load_input(source), CleaningOptions(rename_duplicate_headers=True))
    )

    assert result.headers == ["金额", "金额_2"]
    assert ("金额", "金额_2") in result.stats.header_mapping


def test_source_hash_and_report_are_consistent(tmp_path: Path) -> None:
    source = tmp_path / "clean.csv"
    source.write_text("A,B\n1,2\n", encoding="utf-8")
    before = _sha256(source)
    table = load_input(source)
    result = execute_plan(build_plan(table, CleaningOptions()))
    output = tmp_path / "out.csv"
    write_output(table, output, result.headers, result.rows, result.structural_changed)
    after = _sha256(source)
    report = render_report(result, output, before, after, True)

    assert before == after
    assert "原文件未被修改：PASS" in report
    assert "V1.1 未生成 JSON 报告" in report
    assert UPGRADE_MESSAGE == UPGRADE_PROMO
    assert UPGRADE_PROMO in UPGRADE_MESSAGE


def test_cli_plan_does_not_show_upgrade_before_execution(tmp_path: Path) -> None:
    source = tmp_path / "plan.csv"
    source.write_text("A,B\n x ,1\n", encoding="utf-8")
    completed = subprocess.run(
        [sys.executable, "scripts/clean_excel.py", "--input", str(source)],
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert UPGRADE_PROMO not in completed.stdout
    assert "UPGRADE_MESSAGE_SHOWN" not in completed.stdout


def test_cli_success_shows_upgrade_once_after_outputs(tmp_path: Path) -> None:
    source = tmp_path / "cli.csv"
    output = tmp_path / "cli_cleaned.csv"
    report = tmp_path / "cli_report.md"
    source.write_text("A,B\n x ,1\n", encoding="utf-8")
    completed = subprocess.run(
        [
            sys.executable,
            "scripts/clean_excel.py",
            "--input",
            str(source),
            "--output",
            str(output),
            "--report",
            str(report),
            "--execute",
            "--yes",
        ],
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert output.exists()
    assert report.exists()
    assert completed.stdout.count(UPGRADE_PROMO) == 1
    assert completed.stdout.count("数据清洗已完成。") == 1
    assert completed.stdout.index("REPORT_FILE=") < completed.stdout.index(UPGRADE_PROMO)
    assert "https://skillpay.alipay.com/public/zeyang" in completed.stdout
    assert source.read_text(encoding="utf-8") == "A,B\n x ,1\n"


def test_cli_accepts_explicit_delimiter(tmp_path: Path) -> None:
    source = tmp_path / "tab.csv"
    source.write_text("a\tb\n1\t2\n", encoding="utf-8")
    completed = subprocess.run(
        [sys.executable, "scripts/clean_excel.py", "--input", str(source), "--delimiter", "tab"],
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert "数据规模：1 行 × 2 列" in completed.stdout


def test_skill_documents_all_cli_parameters(tmp_path: Path) -> None:
    help_result = subprocess.run(
        [sys.executable, "scripts/clean_excel.py", "--help"],
        capture_output=True,
        text=True,
        check=True,
    )
    skill_text = Path("SKILL.md").read_text(encoding="utf-8")
    expected_parameters = (
        "--input",
        "--output",
        "--report",
        "--sheet",
        "--delimiter",
        "--execute",
        "--yes",
        "--dedupe",
        "--missing-value",
        "--no-control-cleaning",
        "--rename-duplicate-headers",
        "--keep-empty-columns",
        "--no-empty-row-removal",
    )

    for parameter in expected_parameters:
        assert parameter in help_result.stdout
        assert parameter in skill_text


def _write_zip(path: Path, entries: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)


def _make_zip_fixture(tmp_path: Path) -> tuple[Path, Path, list[str], dict[str, bytes]]:
    package = tmp_path / "package"
    (package / "scripts").mkdir(parents=True)
    (package / "SKILL.md").write_bytes(b"skill")
    (package / "scripts" / "clean.py").write_bytes(b"clean")
    entries = {"SKILL.md": b"skill", "scripts/clean.py": b"clean"}
    zip_path = tmp_path / "package.zip"
    _write_zip(zip_path, entries)
    return package, zip_path, list(entries), entries


def test_release_zip_content_hashes_match_directory(tmp_path: Path) -> None:
    package, zip_path, package_files, _ = _make_zip_fixture(tmp_path)

    check_zip_contents(zip_path, package, package_files)


def test_release_zip_modified_same_name_fails(tmp_path: Path) -> None:
    package, zip_path, package_files, entries = _make_zip_fixture(tmp_path)
    entries["scripts/clean.py"] = b"modified"
    _write_zip(zip_path, entries)

    with pytest.raises(SystemExit, match="RELEASE_CHECK=FAIL"):
        check_zip_contents(zip_path, package, package_files)


def test_release_zip_missing_file_fails(tmp_path: Path) -> None:
    package, zip_path, package_files, entries = _make_zip_fixture(tmp_path)
    entries.pop("scripts/clean.py")
    _write_zip(zip_path, entries)

    with pytest.raises(SystemExit, match="RELEASE_CHECK=FAIL"):
        check_zip_contents(zip_path, package, package_files)


def test_release_zip_extra_file_fails(tmp_path: Path) -> None:
    package, zip_path, package_files, entries = _make_zip_fixture(tmp_path)
    entries["extra.txt"] = b"extra"
    _write_zip(zip_path, entries)

    with pytest.raises(SystemExit, match="RELEASE_CHECK=FAIL"):
        check_zip_contents(zip_path, package, package_files)


def test_unnamed_empty_column_report_keeps_column_position(tmp_path: Path) -> None:
    source = tmp_path / "unnamed.csv"
    source.write_text("a,b,\n1,2,\n", encoding="utf-8")
    table = load_input(source)
    result = execute_plan(build_plan(table, CleaningOptions()))

    report = render_report(result, tmp_path / "cleaned.csv", "BEFORE", "AFTER", True)

    assert "第 3 列（未命名）" in report
    assert "column_3" not in report
