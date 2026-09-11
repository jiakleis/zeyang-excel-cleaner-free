from __future__ import annotations

import csv
import hashlib
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class CleanerError(Exception):
    """Base class for expected user-facing cleaner errors."""


class UnsupportedFormatError(CleanerError):
    pass


class InvalidInputError(CleanerError):
    pass


class SheetSelectionRequired(CleanerError):
    def __init__(self, sheets: list[str]):
        self.sheets = sheets
        super().__init__("XLSX contains multiple non-empty worksheets; specify --sheet: " + ", ".join(sheets))


class UnsafeStructureError(CleanerError):
    pass


DELIMITER_ALIASES = {
    "comma": ",",
    "tab": "\t",
    "semicolon": ";",
    "pipe": "|",
    ",": ",",
    "\\t": "\t",
    ";": ";",
    "|": "|",
}


def normalize_delimiter(value: str | None) -> str | None:
    if value is None:
        return None
    if value in {",", "\t", ";", "|"}:
        return value
    key = value.strip().lower()
    delimiter = DELIMITER_ALIASES.get(key)
    if delimiter is None:
        raise ValueError("delimiter 仅支持 comma、tab、semicolon、pipe、,、\\t、; 或 |。")
    return delimiter


@dataclass
class LoadedTable:
    source_path: Path
    kind: str
    sheet_name: str | None
    headers: list[Any]
    rows: list[list[Any]]
    delimiter: str = ","
    encoding: str = "utf-8"
    workbook: Any | None = None
    worksheet: Any | None = None
    sheet_names: list[str] = field(default_factory=list)
    nonempty_sheet_names: list[str] = field(default_factory=list)
    formula_cells: list[str] = field(default_factory=list)
    merged_ranges: list[str] = field(default_factory=list)
    template_signals: list[str] = field(default_factory=list)

    @property
    def has_template_risk(self) -> bool:
        return bool(self.template_signals)

    @property
    def has_structural_risk(self) -> bool:
        return bool(self.merged_ranges)

    @property
    def width(self) -> int:
        return max([len(self.headers)] + [len(row) for row in self.rows] + [0])

    @property
    def height(self) -> int:
        return len(self.rows)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _is_nonempty(value: Any) -> bool:
    return value is not None and value != ""


def _normalize_width(rows: list[list[Any]], width: int) -> list[list[Any]]:
    return [list(row[:width]) + [None] * max(0, width - len(row)) for row in rows]


def _validate_auto_delimiter(rows: list[list[str]], delimiter: str) -> None:
    if not rows:
        return
    header_width = len(rows[0])
    data_rows = [row for row in rows[1:] if row and any(_is_nonempty(value) for value in row)]
    if not data_rows:
        return

    if any(len(row) != header_width for row in data_rows):
        raise InvalidInputError(
            "CSV 自动分隔符检测存在歧义：表头与数据行列数不一致；"
            "请使用 --delimiter 明确指定 comma、tab、semicolon 或 pipe。"
        )

    blank_headers = [value for value in rows[0] if not _is_nonempty(value)]
    if header_width > 1 and len(blank_headers) > max(1, header_width // 2):
        raise InvalidInputError(
            "CSV 自动分隔符检测存在歧义：出现多个未命名表头；"
            "请使用 --delimiter 明确指定分隔符。"
        )


def _read_csv(path: Path, delimiter: str | None = None) -> LoadedTable:
    raw = path.read_bytes()
    if not raw:
        raise InvalidInputError("CSV 文件为空，无法识别表头和数据区域。")

    last_error: Exception | None = None
    text = None
    encoding = None
    for candidate in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            text = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError as exc:
            last_error = exc
    if text is None or encoding is None:
        raise InvalidInputError(f"CSV 编码无法识别：{last_error}")

    explicit_delimiter = normalize_delimiter(delimiter)
    if explicit_delimiter is None:
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
    else:
        delimiter = explicit_delimiter

    try:
        rows = [
            [value.replace("\r\n", "\n").replace("\r", "\n") for value in row]
            for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
        ]
    except csv.Error as exc:
        raise InvalidInputError(f"CSV 结构无法读取：{exc}") from exc
    if not rows:
        raise InvalidInputError("CSV 文件没有可用行。")
    if explicit_delimiter is None:
        _validate_auto_delimiter(rows, delimiter)

    width = max((len(row) for row in rows), default=0)
    if width == 0:
        raise InvalidInputError("CSV 文件没有可用列。")
    normalized = _normalize_width(rows, width)
    headers = normalized[0]
    data_rows = normalized[1:]
    signals: list[str] = []
    nonempty_columns = [
        index
        for index in range(width)
        if any(_is_nonempty(row[index]) for row in normalized)
    ]
    if nonempty_columns:
        first, last = min(nonempty_columns), max(nonempty_columns)
        internal_empty = [
            index + 1
            for index in range(first, last + 1)
            if not any(_is_nonempty(row[index]) for row in normalized)
        ]
        if internal_empty:
            signals.append("存在数据区域内部的空列，可能承担版式分隔作用")

    return LoadedTable(
        source_path=path,
        kind="csv",
        sheet_name=None,
        headers=headers,
        rows=data_rows,
        delimiter=delimiter,
        encoding=encoding,
        template_signals=signals,
    )


def _worksheet_nonempty(ws: Any) -> bool:
    for row in ws.iter_rows():
        for cell in row:
            if _is_nonempty(cell.value):
                return True
    return False


def _worksheet_bounds(ws: Any) -> tuple[int, int]:
    max_row = 0
    max_col = 0
    for row in ws.iter_rows():
        for cell in row:
            if _is_nonempty(cell.value):
                max_row = max(max_row, cell.row)
                max_col = max(max_col, cell.column)
    return max_row, max_col


def _xlsx_template_signals(ws: Any, max_row: int, max_col: int) -> list[str]:
    signals: list[str] = []
    if ws.merged_cells.ranges:
        signals.append("目标工作表存在合并单元格")
    if any(dimension.hidden for dimension in ws.row_dimensions.values()):
        signals.append("目标工作表存在隐藏行")
    if any(dimension.hidden for dimension in ws.column_dimensions.values()):
        signals.append("目标工作表存在隐藏列")
    nonempty_columns = [
        index
        for index in range(1, max_col + 1)
        if any(_is_nonempty(ws.cell(row, index).value) for row in range(1, max_row + 1))
    ]
    if nonempty_columns:
        first, last = min(nonempty_columns), max(nonempty_columns)
        if any(
            not any(_is_nonempty(ws.cell(row, index).value) for row in range(1, max_row + 1))
            for index in range(first, last + 1)
        ):
            signals.append("数据区域内部存在空列，可能承担版式分隔作用")

    styled_blank = 0
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            if cell.value is None and cell.has_style:
                styled_blank += 1
                if styled_blank >= 3:
                    signals.append("空白单元格带有格式，可能属于模板布局")
                    return list(dict.fromkeys(signals))
    return list(dict.fromkeys(signals))


def _read_xlsx(path: Path, sheet_name: str | None) -> LoadedTable:
    try:
        workbook = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    except Exception as exc:
        raise InvalidInputError(f"XLSX 文件无法打开：{exc}") from exc

    sheet_names = list(workbook.sheetnames)
    nonempty_sheet_names = [name for name in sheet_names if _worksheet_nonempty(workbook[name])]
    if not nonempty_sheet_names:
        raise InvalidInputError("XLSX 工作簿没有可用数据。")
    if sheet_name is None:
        if len(nonempty_sheet_names) > 1:
            raise SheetSelectionRequired(nonempty_sheet_names)
        sheet_name = nonempty_sheet_names[0]
    if sheet_name not in sheet_names:
        raise InvalidInputError(f"找不到工作表：{sheet_name}")

    ws = workbook[sheet_name]
    max_row, max_col = _worksheet_bounds(ws)
    if max_row == 0 or max_col == 0:
        raise InvalidInputError(f"工作表“{sheet_name}”没有可用数据。")
    grid = [
        [ws.cell(row, column).value for column in range(1, max_col + 1)]
        for row in range(1, max_row + 1)
    ]
    headers = grid[0]
    data_rows = grid[1:]
    formula_cells = [
        cell.coordinate
        for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col)
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    merged_ranges = [str(item) for item in ws.merged_cells.ranges]
    return LoadedTable(
        source_path=path,
        kind="xlsx",
        sheet_name=sheet_name,
        headers=headers,
        rows=data_rows,
        workbook=workbook,
        worksheet=ws,
        sheet_names=sheet_names,
        nonempty_sheet_names=nonempty_sheet_names,
        formula_cells=formula_cells,
        merged_ranges=merged_ranges,
        template_signals=_xlsx_template_signals(ws, max_row, max_col),
    )


def load_input(
    path: str | Path,
    sheet_name: str | None = None,
    delimiter: str | None = None,
) -> LoadedTable:
    source = Path(path).expanduser().resolve()
    if not source.exists() or not source.is_file():
        raise InvalidInputError(f"输入文件不存在：{source}")
    if source.stat().st_size == 0:
        raise InvalidInputError("输入文件为空。")
    suffix = source.suffix.lower()
    if suffix == ".csv":
        return _read_csv(source, delimiter)
    if suffix == ".xlsx":
        if delimiter is not None:
            raise InvalidInputError("--delimiter 仅适用于 CSV 文件。")
        return _read_xlsx(source, sheet_name)
    raise UnsupportedFormatError(f"不支持的文件格式：{suffix or '(无扩展名)'}；V1 仅支持 .xlsx 和 .csv。")


def write_csv(path: Path, headers: list[Any], rows: list[list[Any]], delimiter: str) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=delimiter, lineterminator="\n")
        writer.writerow(["" if value is None else value for value in headers])
        writer.writerows([["" if value is None else value for value in row] for row in rows])


def write_xlsx(
    table: LoadedTable,
    path: Path,
    headers: list[Any],
    rows: list[list[Any]],
    structural_changed: bool,
    source_row_indices: list[int] | None = None,
    source_column_indices: list[int] | None = None,
) -> None:
    if table.workbook is None or table.worksheet is None:
        raise InvalidInputError("XLSX 工作簿上下文丢失，无法安全输出。")
    workbook = table.workbook
    ws = table.worksheet
    if structural_changed:
        if table.formula_cells or table.merged_ranges:
            raise UnsafeStructureError("目标工作表含公式或合并单元格，不能执行会改变行列结构的操作。")
        row_indices = source_row_indices if source_row_indices is not None else list(range(len(rows)))
        column_indices = (
            source_column_indices
            if source_column_indices is not None
            else list(range(len(headers)))
        )
        number_formats: dict[tuple[int, int], str] = {}
        for output_row, source_row in enumerate(row_indices, start=2):
            for output_column, source_column in enumerate(column_indices, start=1):
                source_cell = ws.cell(source_row + 2, source_column + 1)
                number_formats[(output_row, output_column)] = source_cell.number_format
        ws.delete_rows(1, ws.max_row)
        if ws.max_column:
            ws.delete_cols(1, ws.max_column)
        ws.append(list(headers))
        for row_index, row in enumerate(rows, start=2):
            ws.append(list(row))
            for column_index in range(1, len(row) + 1):
                ws.cell(row_index, column_index).number_format = number_formats.get(
                    (row_index, column_index), "General"
                )
    else:
        for row_index, row in enumerate([headers] + rows, start=1):
            for column_index, value in enumerate(row, start=1):
                current = ws.cell(row_index, column_index).value
                if isinstance(current, str) and current.startswith("="):
                    continue
                ws.cell(row_index, column_index).value = value
    workbook.save(path)


def write_output(
    table: LoadedTable,
    path: str | Path,
    headers: list[Any],
    rows: list[list[Any]],
    structural_changed: bool,
    source_row_indices: list[int] | None = None,
    source_column_indices: list[int] | None = None,
) -> None:
    output = Path(path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        raise InvalidInputError(f"输出文件已存在，为避免覆盖未写入：{output}")
    if table.kind == "csv":
        write_csv(output, headers, rows, table.delimiter)
    else:
        write_xlsx(
            table,
            output,
            headers,
            rows,
            structural_changed,
            source_row_indices,
            source_column_indices,
        )
